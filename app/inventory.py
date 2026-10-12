import csv
import io
import os
import time

import psycopg
from openpyxl import load_workbook
from packageurl import PackageURL
from psycopg.rows import dict_row

COLUMNS = (
    "Project",
    "PURL",
    "Business Process",
    "Internet-facing?",
    "Criticality",
    "Potential Consequence",
)

DB_FIELDS = {
    "Project": "project",
    "PURL": "purl",
    "Business Process": "business_process",
    "Internet-facing?": "internet_facing",
    "Criticality": "criticality",
    "Potential Consequence": "potential_consequence",
}

CANONICAL_COLUMNS = {column.casefold(): column for column in COLUMNS}

MAX_UPLOAD_BYTES = 10 * 1024 * 1024

CREATE_TABLE = """
CREATE TABLE IF NOT EXISTS assets (
    id INTEGER GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    project TEXT NOT NULL,
    purl TEXT NOT NULL,
    business_process TEXT NOT NULL,
    internet_facing TEXT NOT NULL,
    criticality TEXT NOT NULL,
    potential_consequence TEXT NOT NULL
)
"""

ADD_PROJECT = "ALTER TABLE assets ADD COLUMN IF NOT EXISTS project TEXT"

CREATE_META = """
CREATE TABLE IF NOT EXISTS inventory_meta (
    id INTEGER PRIMARY KEY CHECK (id = 1),
    extracted_at TIMESTAMPTZ NOT NULL
)
"""


def database_url() -> str:
    url = os.environ.get("DATABASE_URL")
    if not url:
        raise RuntimeError("DATABASE_URL is not set")
    return url


def connect():
    return psycopg.connect(database_url(), row_factory=dict_row)


def init_db() -> None:
    last_error = None
    for _ in range(30):
        try:
            with connect() as connection:
                connection.execute(CREATE_TABLE)
                connection.execute(ADD_PROJECT)
                _rename_business_purpose(connection)
                connection.execute(CREATE_META)
                connection.commit()
            return
        except psycopg.OperationalError as exc:
            last_error = exc
            time.sleep(1)
    raise last_error


def _rename_business_purpose(connection) -> None:
    columns = {
        row["column_name"]
        for row in connection.execute(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = 'public' AND table_name = 'assets'
            """
        ).fetchall()
    }
    if "business_purpose" not in columns:
        return
    if "business_process" in columns:
        connection.execute("ALTER TABLE assets DROP COLUMN business_process")
    connection.execute("ALTER TABLE assets RENAME COLUMN business_purpose TO business_process")


def list_inventory() -> dict:
    with connect() as connection:
        rows = connection.execute(
            """
            SELECT project, purl, business_process, internet_facing, criticality, potential_consequence
            FROM assets
            ORDER BY id
            """
        ).fetchall()
        meta = connection.execute(
            "SELECT extracted_at FROM inventory_meta WHERE id = 1"
        ).fetchone()
    extracted_at = meta["extracted_at"].isoformat() if meta else None
    return {"items": [dict(row) for row in rows], "extracted_at": extracted_at}


def replace_inventory(rows: list[dict]) -> None:
    with connect() as connection:
        with connection.transaction():
            connection.execute("DELETE FROM assets")
            with connection.cursor() as cursor:
                cursor.executemany(
                    """
                    INSERT INTO assets (
                        project, purl, business_process, internet_facing, criticality, potential_consequence
                    )
                    VALUES (
                        %(project)s, %(purl)s, %(business_process)s, %(internet_facing)s,
                        %(criticality)s, %(potential_consequence)s
                    )
                    """,
                    rows,
                )
            connection.execute(
                """
                INSERT INTO inventory_meta (id, extracted_at)
                VALUES (1, NOW())
                ON CONFLICT (id) DO UPDATE SET extracted_at = EXCLUDED.extracted_at
                """
            )


def process_upload(filename: str, content: bytes) -> tuple[int, dict]:
    if len(content) > MAX_UPLOAD_BYTES:
        return 422, {"ok": False, "error": "The file is larger than 10 MB."}
    if not content:
        return 422, {"ok": False, "error": "The file is empty."}

    parsed, error = _read_rows(filename, content)
    if error:
        return 422, {"ok": False, "error": error}

    review = []
    stored = []
    any_invalid = False
    for row_number, values in parsed:
        errors = _validate_row(values)
        if errors:
            any_invalid = True
        review.append({"row": row_number, "values": values, "errors": errors})
        stored.append({DB_FIELDS[column]: values[column] for column in COLUMNS})

    if any_invalid:
        return 422, {"ok": False, "rows": review}

    replace_inventory(stored)
    return 200, {"ok": True, "count": len(stored)}


def _read_rows(filename: str, content: bytes) -> tuple[list[tuple[int, dict]] | None, str | None]:
    lowered = filename.lower()
    try:
        if lowered.endswith(".csv"):
            table = _read_csv(content)
        elif lowered.endswith(".xlsx"):
            table = _read_xlsx(content)
        else:
            return None, "Upload a .csv or .xlsx file."
    except ValueError as exc:
        return None, str(exc)
    except Exception:
        return None, "The file could not be read."

    if not table:
        return None, "The file has no header row."

    header = _trim_trailing_empty(table[0])
    header_error = _header_error(header)
    if header_error:
        return None, header_error

    positions = {}
    for index, name in enumerate(header):
        canonical = CANONICAL_COLUMNS.get(name.casefold())
        if canonical:
            positions[canonical] = index
    rows = []
    for offset, cells in enumerate(table[1:], start=2):
        values = {column: _cell(cells, positions[column]) for column in COLUMNS}
        if any(values.values()):
            rows.append((offset, values))

    if not rows:
        return None, "The file has no data rows."
    return rows, None


def _header_error(header: list[str]) -> str | None:
    if any(not name for name in header):
        return "The header row has an empty column name."

    folded = [name.casefold() for name in header]
    if len(folded) != len(set(folded)):
        return "The header row contains duplicate column names."

    present = set(folded)
    missing = [column for column in COLUMNS if column.casefold() not in present]
    unexpected = [column for column in header if column.casefold() not in CANONICAL_COLUMNS]
    if not missing and not unexpected:
        return None

    parts = []
    if missing:
        parts.append("Missing columns: " + ", ".join(missing))
    if unexpected:
        parts.append("Unexpected columns: " + ", ".join(unexpected))
    return ". ".join(parts)


def _validate_row(values: dict) -> dict:
    errors = {}

    if not values["Project"]:
        errors["Project"] = "must not be empty"

    purl = values["PURL"]
    if not purl:
        errors["PURL"] = "not a valid PURL"
    else:
        try:
            PackageURL.from_string(purl)
        except ValueError:
            errors["PURL"] = "not a valid PURL"

    if not values["Business Process"]:
        errors["Business Process"] = "must not be empty"

    if values["Internet-facing?"] not in {"Yes", "No"}:
        errors["Internet-facing?"] = 'must be "Yes" or "No"'

    if values["Criticality"] not in {"Low", "Medium", "High"}:
        errors["Criticality"] = 'must be "Low", "Medium", or "High"'

    if not values["Potential Consequence"]:
        errors["Potential Consequence"] = "must not be empty"

    return errors


def _trim_trailing_empty(cells: list[str]) -> list[str]:
    trimmed = list(cells)
    while trimmed and not trimmed[-1].strip():
        trimmed.pop()
    return trimmed


def _cell(cells: list[str], index: int) -> str:
    if index >= len(cells):
        return ""
    return cells[index].strip()


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("The CSV file must be UTF-8.") from exc
    return [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]


def _read_xlsx(content: bytes) -> list[list[str]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        table = []
        for row in worksheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cell.strip() for cell in cells):
                table.append(cells)
        return table
    finally:
        workbook.close()
