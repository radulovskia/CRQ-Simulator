import csv
import io
import os
import re
import time

import psycopg
from openpyxl import load_workbook
from psycopg.rows import dict_row

MAX_UPLOAD_BYTES = 10 * 1024 * 1024

COLUMN_LABELS = {
    "OU ID": "Organisational Unit ID",
    "BP ID": "Business Process ID",
    "RTO (hours)": "Recovery Time Objective (hours)",
    "Downtime Cost per Hour before RTO": "Downtime Cost per Hour before Recovery Time Objective",
    "Downtime Cost per Hour post RTO": "Downtime Cost per Hour post Recovery Time Objective",
    "FTE Count": "Full-Time Equivalent Count",
    "Business Owner OU ID": "Business Owner Organisational Unit ID",
    "Technical Custodian OU ID": "Technical Custodian Organisational Unit ID",
    "Owner OU ID": "Owner Organisational Unit ID",
    "Risk Owner OU ID": "Risk Owner Organisational Unit ID",
    "Base TEF Min": "Base Threat Event Frequency Min",
    "Base TEF Mode": "Base Threat Event Frequency Mode",
    "Base TEF Max": "Base Threat Event Frequency Max",
    "TEF Min": "Threat Event Frequency Min",
    "TEF Mode": "Threat Event Frequency Mode",
    "TEF Max": "Threat Event Frequency Max",
    "TEF Basis": "Threat Event Frequency Basis",
    "PURL": "Package URL",
}


def database_url() -> str:
    url = os.environ.get("DATABASE_URL")
    if not url:
        raise RuntimeError("DATABASE_URL is not set")
    return url


def connect():
    return psycopg.connect(database_url(), row_factory=dict_row)


def rename_column(connection, table: str, old: str, new: str) -> None:
    found = connection.execute(
        """
        SELECT 1
        FROM information_schema.columns
        WHERE table_schema = 'public' AND table_name = %s AND column_name = %s
        """,
        (table, old),
    ).fetchone()
    if not found:
        return
    connection.execute(f"ALTER TABLE {table} RENAME COLUMN {old} TO {new}")


def wait_for_database(apply) -> None:
    last_error = None
    for _ in range(30):
        try:
            with connect() as connection:
                apply(connection)
                connection.commit()
            return
        except psycopg.OperationalError as exc:
            last_error = exc
            time.sleep(1)
    raise last_error


def read_tabular(filename: str, content: bytes, columns: tuple[str, ...]):
    if len(content) > MAX_UPLOAD_BYTES:
        return None, "The file is larger than 10 MB."
    if not content:
        return None, "The file is empty."

    canonical = {column.casefold(): column for column in columns}
    for official, label in COLUMN_LABELS.items():
        if official in columns:
            canonical[label.casefold()] = official
    lowered = filename.lower()
    try:
        if lowered.endswith(".csv"):
            table = _read_csv(content)
        elif lowered.endswith(".xlsx"):
            table = _read_xlsx(content)
        else:
            return None, "Upload a .csv or .xlsx file."
    except ValueError as exc:
        return None, str(exc)
    except Exception:
        return None, "The file could not be read."

    if not table:
        return None, "The file has no header row."

    header = _trim_trailing_empty(table[0])
    header_error = _header_error(header, columns, canonical)
    if header_error:
        return None, header_error

    positions = {}
    for index, name in enumerate(header):
        column = canonical.get(name.casefold())
        if column:
            positions[column] = index
    rows = []
    for offset, cells in enumerate(table[1:], start=2):
        values = {column: _cell(cells, positions[column]) for column in columns}
        if any(values.values()):
            rows.append((offset, values))
    if not rows:
        return None, "The file has no data rows."
    return rows, None


def split_ids(value: str) -> list[str]:
    return [part.strip() for part in str(value or "").split(",") if part.strip()]


def duplicate_ids(values: list[str], label: str) -> str | None:
    seen = set()
    duplicates = []
    for value in values:
        if not value:
            continue
        if value in seen and value not in duplicates:
            duplicates.append(value)
        seen.add(value)
    if not duplicates:
        return None
    return f"The file contains duplicate {label}: " + ", ".join(duplicates)


def _header_error(header: list[str], columns: tuple[str, ...], canonical: dict[str, str]) -> str | None:
    if any(not name for name in header):
        return "The header row has an empty column name."

    folded = [name.casefold() for name in header]
    resolved = [canonical[name] for name in folded if name in canonical]
    if len(folded) != len(set(folded)) or len(resolved) != len(set(resolved)):
        return "The header row contains duplicate column names."

    present = set(resolved)
    missing = [column for column in columns if column not in present]
    unexpected = [column for column in header if column.casefold() not in canonical]
    if not missing and not unexpected:
        return None

    parts = []
    if missing:
        parts.append("Missing columns: " + ", ".join(COLUMN_LABELS.get(column, column) for column in missing))
    if unexpected:
        parts.append("Unexpected columns: " + ", ".join(unexpected))
    return ". ".join(parts)


def _trim_trailing_empty(cells: list[str]) -> list[str]:
    trimmed = list(cells)
    while trimmed and not trimmed[-1].strip():
        trimmed.pop()
    return trimmed


def _cell(cells: list[str], index: int) -> str:
    if index >= len(cells):
        return ""
    return cells[index].strip()


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("The CSV file must be UTF-8.") from exc
    return [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]


GROUP_ID = re.compile(r"^G\d{4}$")
MITIGATION_ID = re.compile(r"^M\d{4}$")
TECHNIQUE_ID = re.compile(r"^T\d{4}(\.\d{3})?$")


def parse_number(text) -> float | None:
    cleaned = str(text if text is not None else "").strip()
    if cleaned == "":
        return None
    try:
        value = float(cleaned)
    except ValueError as exc:
        raise ValueError("must be a number") from exc
    if value != value or value in (float("inf"), float("-inf")):
        raise ValueError("must be a number")
    return value


def parse_int(text) -> int | None:
    value = parse_number(text)
    if value is None:
        return None
    if not float(value).is_integer():
        raise ValueError("must be a whole number")
    return int(value)


def number_error(
    text,
    *,
    allow_blank: bool,
    minimum: float | None = None,
    maximum: float | None = None,
    exclusive_max: bool = False,
    integer: bool = False,
) -> str | None:
    cleaned = str(text if text is not None else "").strip()
    if cleaned == "":
        return None if allow_blank else "must not be empty"
    try:
        value = parse_int(cleaned) if integer else parse_number(cleaned)
    except ValueError as exc:
        return str(exc)
    if value is None:
        return None if allow_blank else "must not be empty"
    if minimum is not None and value < minimum:
        return f"must be >= {minimum}"
    if maximum is not None:
        if exclusive_max and value >= maximum:
            return f"must be < {maximum}"
        if not exclusive_max and value > maximum:
            return f"must be <= {maximum}"
    return None


def triple_error(values: dict, prefix: str) -> dict[str, str]:
    errors = {}
    parsed = {}
    for part in ("Min", "Mode", "Max"):
        column = f"{prefix} {part}"
        error = number_error(values.get(column, ""), allow_blank=False, minimum=0)
        if error:
            errors[column] = error
        else:
            parsed[part] = parse_number(values.get(column, ""))
    if len(parsed) == 3 and not (parsed["Min"] <= parsed["Mode"] <= parsed["Max"]):
        errors[f"{prefix} Mode"] = "must satisfy Min <= Mode <= Max"
    return errors


def enum_error(value: str, options: tuple[str, ...]) -> str | None:
    if value in options:
        return None
    if len(options) == 2:
        return f'must be "{options[0]}" or "{options[1]}"'
    head = ", ".join(f'"{option}"' for option in options[:-1])
    return f'must be {head}, or "{options[-1]}"'


def attack_id_error(text: str, pattern: re.Pattern, label: str) -> str | None:
    raw = str(text or "").strip()
    if not raw:
        return None
    parts = [part.strip() for part in raw.split(",")]
    if any(not part for part in parts):
        return f"must contain only {label}"
    if any(pattern.fullmatch(part) is None for part in parts):
        return f"must contain only {label}"
    if len(parts) != len(set(parts)):
        return f"contains duplicate {label}"
    return None


def normalise_id_list(text: str) -> str:
    return ", ".join(part.strip() for part in str(text or "").split(",") if part.strip())


def technique_chain_error(text: str) -> str | None:
    raw = str(text or "").strip()
    if not raw:
        return "must not be empty"
    tokens = [token.strip() for token in re.split(r"->|→|,", raw)]
    if not tokens or any(not token for token in tokens):
        return "must contain ATT&CK technique IDs"
    if any(TECHNIQUE_ID.fullmatch(token) is None for token in tokens):
        return "must contain ATT&CK technique IDs"
    return None


def normalise_technique_chain(text: str) -> str:
    tokens = [token.strip() for token in re.split(r"->|→|,", str(text or "")) if token.strip()]
    return " -> ".join(tokens)


def _read_xlsx(content: bytes) -> list[list[str]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        table = []
        for row in worksheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cell.strip() for cell in cells):
                table.append(cells)
        return table
    finally:
        workbook.close()
