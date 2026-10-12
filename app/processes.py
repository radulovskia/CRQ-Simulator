import csv
import io
import os
import time
from datetime import datetime, timezone

import psycopg
from openpyxl import load_workbook
from psycopg.rows import dict_row

COLUMNS = (
    "Business Process",
    "Organisational Unit",
    "Criticality",
)

DB_FIELDS = {
    "Business Process": "business_process",
    "Organisational Unit": "organisational_unit",
    "Criticality": "criticality",
}

CANONICAL_COLUMNS = {column.casefold(): column for column in COLUMNS}
CRITICALITY = {"Low", "Medium", "High"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024

CREATE_TABLE = """
CREATE TABLE IF NOT EXISTS business_processes (
    id INTEGER GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    business_process TEXT NOT NULL,
    organisational_unit TEXT NOT NULL,
    criticality TEXT NOT NULL,
    date_added TIMESTAMPTZ NOT NULL
)
"""

CREATE_META = """
CREATE TABLE IF NOT EXISTS business_process_meta (
    id INTEGER PRIMARY KEY CHECK (id = 1),
    extracted_at TIMESTAMPTZ NOT NULL
)
"""

TOUCH_META = """
INSERT INTO business_process_meta (id, extracted_at)
VALUES (1, NOW())
ON CONFLICT (id) DO UPDATE SET extracted_at = EXCLUDED.extracted_at
"""


def database_url() -> str:
    url = os.environ.get("DATABASE_URL")
    if not url:
        raise RuntimeError("DATABASE_URL is not set")
    return url


def connect():
    return psycopg.connect(database_url(), row_factory=dict_row)


def init_processes() -> None:
    last_error = None
    for _ in range(30):
        try:
            with connect() as connection:
                connection.execute(CREATE_TABLE)
                connection.execute(CREATE_META)
                connection.commit()
            return
        except psycopg.OperationalError as exc:
            last_error = exc
            time.sleep(1)
    raise last_error


def list_processes() -> dict:
    with connect() as connection:
        rows = connection.execute(
            """
            SELECT id, business_process, organisational_unit, criticality, date_added
            FROM business_processes
            ORDER BY id
            """
        ).fetchall()
        meta = connection.execute(
            "SELECT extracted_at FROM business_process_meta WHERE id = 1"
        ).fetchone()
    extracted_at = meta["extracted_at"].isoformat() if meta else None
    return {
        "items": [_public_row(row) for row in rows],
        "extracted_at": extracted_at,
    }


def missing_names(names: list[str]) -> list[str]:
    with connect() as connection:
        existing = {
            row["business_process"].casefold()
            for row in connection.execute("SELECT business_process FROM business_processes").fetchall()
        }
    missing = []
    seen = set()
    for name in names:
        cleaned = name.strip()
        if not cleaned:
            continue
        key = cleaned.casefold()
        if key in seen or key in existing:
            continue
        seen.add(key)
        missing.append(cleaned)
    return missing


def process_catalog_upload(filename: str, content: bytes) -> tuple[int, dict]:
    if len(content) > MAX_UPLOAD_BYTES:
        return 422, {"ok": False, "error": "The file is larger than 10 MB."}
    if not content:
        return 422, {"ok": False, "error": "The file is empty."}

    parsed, error = _read_rows(filename, content)
    if error:
        return 422, {"ok": False, "error": error}

    review = []
    stored = []
    any_invalid = False
    for row_number, values in parsed:
        errors = _validate_values(values)
        if errors:
            any_invalid = True
        review.append({"row": row_number, "values": values, "errors": errors})
        stored.append({DB_FIELDS[column]: values[column] for column in COLUMNS})

    if any_invalid:
        return 422, {"ok": False, "rows": review}

    _replace_catalog(stored)
    stored_rows = list_processes()
    return 200, {"ok": True, "count": len(stored_rows["items"]), **stored_rows}


def add_entry(fields: dict) -> tuple[int, dict]:
    errors = _validate_fields(fields)
    if errors:
        return 422, {"ok": False, "errors": errors}
    _insert_entries([_clean_fields(fields)])
    return 200, {"ok": True, **list_processes()}


def add_entries(entries: list[dict]) -> tuple[int, dict]:
    if not entries:
        return 422, {"ok": False, "error": "No business processes were provided."}
    problems = []
    cleaned = []
    for index, fields in enumerate(entries):
        errors = _validate_fields(fields)
        if errors:
            problems.append({"index": index, "errors": errors})
        else:
            cleaned.append(_clean_fields(fields))
    if problems:
        return 422, {"ok": False, "errors": problems}
    _insert_entries(cleaned)
    return 200, {"ok": True, "count": len(cleaned), **list_processes()}


def update_entry(entry_id: int, fields: dict) -> tuple[int, dict]:
    errors = _validate_fields(fields)
    if errors:
        return 422, {"ok": False, "errors": errors}
    cleaned = _clean_fields(fields)
    with connect() as connection:
        with connection.transaction():
            updated = connection.execute(
                """
                UPDATE business_processes
                SET business_process = %(business_process)s,
                    organisational_unit = %(organisational_unit)s,
                    criticality = %(criticality)s
                WHERE id = %(id)s
                RETURNING id
                """,
                {**cleaned, "id": entry_id},
            ).fetchone()
            if not updated:
                return 404, {"ok": False, "error": "That business process was not found."}
            connection.execute(TOUCH_META)
    return 200, {"ok": True, **list_processes()}


def _public_row(row: dict) -> dict:
    return {
        "id": row["id"],
        "business_process": row["business_process"],
        "organisational_unit": row["organisational_unit"],
        "criticality": row["criticality"],
        "date_added": row["date_added"].isoformat() if row["date_added"] else None,
    }


def _replace_catalog(rows: list[dict]) -> None:
    now = datetime.now(timezone.utc)
    with connect() as connection:
        with connection.transaction():
            existing = {
                row["business_process"].casefold(): row["date_added"]
                for row in connection.execute(
                    "SELECT business_process, date_added FROM business_processes"
                ).fetchall()
            }
            connection.execute("DELETE FROM business_processes")
            if rows:
                stamped = [
                    {**row, "date_added": existing.get(row["business_process"].casefold(), now)}
                    for row in rows
                ]
                with connection.cursor() as cursor:
                    cursor.executemany(
                        """
                        INSERT INTO business_processes (
                            business_process, organisational_unit, criticality, date_added
                        )
                        VALUES (
                            %(business_process)s, %(organisational_unit)s, %(criticality)s, %(date_added)s
                        )
                        """,
                        stamped,
                    )
            connection.execute(TOUCH_META)


def _insert_entries(rows: list[dict]) -> None:
    with connect() as connection:
        with connection.transaction():
            with connection.cursor() as cursor:
                cursor.executemany(
                    """
                    INSERT INTO business_processes (
                        business_process, organisational_unit, criticality, date_added
                    )
                    VALUES (
                        %(business_process)s, %(organisational_unit)s, %(criticality)s, NOW()
                    )
                    """,
                    rows,
                )
            connection.execute(TOUCH_META)


def _clean_fields(fields: dict) -> dict:
    return {
        "business_process": str(fields.get("business_process") or "").strip(),
        "organisational_unit": str(fields.get("organisational_unit") or "").strip(),
        "criticality": str(fields.get("criticality") or "").strip(),
    }


def _validate_fields(fields: dict) -> dict:
    cleaned = _clean_fields(fields)
    return _validate_values(
        {
            "Business Process": cleaned["business_process"],
            "Organisational Unit": cleaned["organisational_unit"],
            "Criticality": cleaned["criticality"],
        }
    )


def _validate_values(values: dict) -> dict:
    errors = {}
    if not values["Business Process"]:
        errors["Business Process"] = "must not be empty"
    if not values["Organisational Unit"]:
        errors["Organisational Unit"] = "must not be empty"
    if values["Criticality"] not in CRITICALITY:
        errors["Criticality"] = 'must be "Low", "Medium", or "High"'
    return errors


def _read_rows(filename: str, content: bytes) -> tuple[list[tuple[int, dict]] | None, str | None]:
    lowered = filename.lower()
    try:
        if lowered.endswith(".csv"):
            table = _read_csv(content)
        elif lowered.endswith(".xlsx"):
            table = _read_xlsx(content)
        else:
            return None, "Upload a .csv or .xlsx file."
    except ValueError as exc:
        return None, str(exc)
    except Exception:
        return None, "The file could not be read."

    if not table:
        return None, "The file has no header row."

    header = _trim_trailing_empty(table[0])
    header_error = _header_error(header)
    if header_error:
        return None, header_error

    positions = {}
    for index, name in enumerate(header):
        canonical = CANONICAL_COLUMNS.get(name.casefold())
        if canonical:
            positions[canonical] = index
    rows = []
    for offset, cells in enumerate(table[1:], start=2):
        values = {column: _cell(cells, positions[column]) for column in COLUMNS}
        if any(values.values()):
            rows.append((offset, values))

    if not rows:
        return None, "The file has no data rows."
    return rows, None


def _header_error(header: list[str]) -> str | None:
    if any(not name for name in header):
        return "The header row has an empty column name."

    folded = [name.casefold() for name in header]
    if len(folded) != len(set(folded)):
        return "The header row contains duplicate column names."

    present = set(folded)
    missing = [column for column in COLUMNS if column.casefold() not in present]
    unexpected = [column for column in header if column.casefold() not in CANONICAL_COLUMNS]
    if not missing and not unexpected:
        return None

    parts = []
    if missing:
        parts.append("Missing columns: " + ", ".join(missing))
    if unexpected:
        parts.append("Unexpected columns: " + ", ".join(unexpected))
    return ". ".join(parts)


def _trim_trailing_empty(cells: list[str]) -> list[str]:
    trimmed = list(cells)
    while trimmed and not trimmed[-1].strip():
        trimmed.pop()
    return trimmed


def _cell(cells: list[str], index: int) -> str:
    if index >= len(cells):
        return ""
    return cells[index].strip()


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("The CSV file must be UTF-8.") from exc
    return [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]


def _read_xlsx(content: bytes) -> list[list[str]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        table = []
        for row in worksheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cell.strip() for cell in cells):
                table.append(cells)
        return table
    finally:
        workbook.close()
