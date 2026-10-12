import csv
import io
import os
import time

import psycopg
from openpyxl import load_workbook
from psycopg.rows import dict_row

MAX_UPLOAD_BYTES = 10 * 1024 * 1024


def database_url() -> str:
    url = os.environ.get("DATABASE_URL")
    if not url:
        raise RuntimeError("DATABASE_URL is not set")
    return url


def connect():
    return psycopg.connect(database_url(), row_factory=dict_row)


def wait_for_database(apply) -> None:
    last_error = None
    for _ in range(30):
        try:
            with connect() as connection:
                apply(connection)
                connection.commit()
            return
        except psycopg.OperationalError as exc:
            last_error = exc
            time.sleep(1)
    raise last_error


def read_tabular(filename: str, content: bytes, columns: tuple[str, ...]):
    if len(content) > MAX_UPLOAD_BYTES:
        return None, "The file is larger than 10 MB."
    if not content:
        return None, "The file is empty."

    canonical = {column.casefold(): column for column in columns}
    lowered = filename.lower()
    try:
        if lowered.endswith(".csv"):
            table = _read_csv(content)
        elif lowered.endswith(".xlsx"):
            table = _read_xlsx(content)
        else:
            return None, "Upload a .csv or .xlsx file."
    except ValueError as exc:
        return None, str(exc)
    except Exception:
        return None, "The file could not be read."

    if not table:
        return None, "The file has no header row."

    header = _trim_trailing_empty(table[0])
    header_error = _header_error(header, columns, canonical)
    if header_error:
        return None, header_error

    positions = {}
    for index, name in enumerate(header):
        column = canonical.get(name.casefold())
        if column:
            positions[column] = index
    rows = []
    for offset, cells in enumerate(table[1:], start=2):
        values = {column: _cell(cells, positions[column]) for column in columns}
        if any(values.values()):
            rows.append((offset, values))
    if not rows:
        return None, "The file has no data rows."
    return rows, None


def split_ids(value: str) -> list[str]:
    return [part.strip() for part in str(value or "").split(",") if part.strip()]


def duplicate_ids(values: list[str], label: str) -> str | None:
    seen = set()
    duplicates = []
    for value in values:
        if not value:
            continue
        if value in seen and value not in duplicates:
            duplicates.append(value)
        seen.add(value)
    if not duplicates:
        return None
    return f"The file contains duplicate {label}: " + ", ".join(duplicates)


def _header_error(header: list[str], columns: tuple[str, ...], canonical: dict[str, str]) -> str | None:
    if any(not name for name in header):
        return "The header row has an empty column name."

    folded = [name.casefold() for name in header]
    if len(folded) != len(set(folded)):
        return "The header row contains duplicate column names."

    present = set(folded)
    missing = [column for column in columns if column.casefold() not in present]
    unexpected = [column for column in header if column.casefold() not in canonical]
    if not missing and not unexpected:
        return None

    parts = []
    if missing:
        parts.append("Missing columns: " + ", ".join(missing))
    if unexpected:
        parts.append("Unexpected columns: " + ", ".join(unexpected))
    return ". ".join(parts)


def _trim_trailing_empty(cells: list[str]) -> list[str]:
    trimmed = list(cells)
    while trimmed and not trimmed[-1].strip():
        trimmed.pop()
    return trimmed


def _cell(cells: list[str], index: int) -> str:
    if index >= len(cells):
        return ""
    return cells[index].strip()


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("The CSV file must be UTF-8.") from exc
    return [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]


def _read_xlsx(content: bytes) -> list[list[str]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        table = []
        for row in worksheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cell.strip() for cell in cells):
                table.append(cells)
        return table
    finally:
        workbook.close()
